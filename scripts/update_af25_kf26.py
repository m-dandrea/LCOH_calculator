"""Add comparable AF25/KF26 inputs after the workbook and catalogue imports.

The published price bases are retained. No CPI series was supplied, so 2025/2026
prices are not silently rebased to the original workbook's 2019 price level.
"""
import json
from pathlib import Path
from statistics import mean
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
UPLOAD = ROOT.parent / 'upload'
DATA = ROOT / 'data' / 'workbook.json'
SOURCES = ROOT / 'data' / 'input-sources.json'
data = json.loads(DATA.read_text(encoding='utf-8'))
sources = json.loads(SOURCES.read_text(encoding='utf-8'))
af_name = 'offentligt_dataset_AF25_20251107_v2.xlsx'
kf_name = 'KF26 - Resultater - Tal bag figurer september 2026.xlsx'
af = openpyxl.load_workbook(UPLOAD / af_name, read_only=True, data_only=True)
kf = openpyxl.load_workbook(UPLOAD / kf_name, read_only=True, data_only=True)

def source(category, identifier, field, filename, cell, parameter, unit, basis, conversion=1):
    sources.setdefault(category, {}).setdefault(str(identifier), {})[field] = {
        'file': filename, 'cell': cell, 'parameter': parameter,
        'unit': unit, 'priceYear': basis, 'conversion': conversion,
    }

# The original Fuel price projection is explicitly CIF/import prices, in DKK/GJ.
# Industrial wood pellets (column J) match AF25's imported wood pellets.
# AF25 has no matching imported diesel/gasoil series: preserve that column.
fuel = af['Brændselspriser']
for record in data['fuelPrices']:
    year = record['year']
    if not 2025 <= year <= 2050:
        continue
    col = year - 2025 + 3
    letter = openpyxl.utils.get_column_letter(col)
    for field, row, label in [('gas', 8, 'Natural gas import'),
                              ('coal', 6, 'Coal import'),
                              ('wood', 27, 'Industrial wood pellet import')]:
        value = fuel.cell(row, col).value
        if not isinstance(value, (int, float)):
            raise ValueError(f'Missing AF25 {label} {year}')
        record[field] = value
        source('fuelPrices', year, field, af_name, f'Brændselspriser!{letter}{row}',
               label, 'DKK/GJ', 2025)

# An ETS1 allowance price is an explicit scenario, rather than a replacement for
# the workbook's general carbon-tax options. The model expects EUR/t.
quota = af['CO2-kvotepris']
for year in [2030, 2035, 2040, 2050]:
    col = year - 2025 + 3
    letter = openpyxl.utils.get_column_letter(col)
    name = f'AF25 ETS1 allowance - {year}'
    value = quota.cell(4, col).value
    if not isinstance(value, (int, float)):
        raise ValueError(f'Missing AF25 ETS1 {year}')
    item = {'name': name, 'year': year, 'price': value / data['defaults']['dkkPerEur']}
    data['carbon'] = [r for r in data['carbon'] if r['name'] != name] + [item]
    source('carbon', name, 'price', af_name, f'CO2-kvotepris!{letter}4',
           'ETS1 allowance price', 'DKK/t CO₂', 2025,
           f"divide by workbook DKK/EUR ({data['defaults']['dkkPerEur']})")

# KF26 publishes 8,760 hourly *spot* prices for 2030 and 2035. The simple
# annual mean is a transparent reference price; an electrolyser's hourly dispatch
# or power purchase agreement would produce a different realised price.
spot = kf['25.1']
hourly = list(spot.iter_rows(min_row=4, max_row=8763, min_col=3, max_col=4, values_only=True))
for year, col in [(2030, 3), (2035, 4)]:
    prices = [row[col - 3] for row in hourly]
    if len(prices) != 8760 or not all(isinstance(x, (int, float)) for x in prices):
        raise ValueError(f'Incomplete KF26 hourly spot series for {year}')
    name = f'KF26 grid spot average - {year}'
    # Tariff and grid emissions have no direct match in these supplied datasets.
    old_grid = next(r for r in data['electricity'] if r['name'] == 'El grid, no PtX nor island')
    item = {'name': name, 'price': mean(prices) / data['defaults']['dkkPerEur'],
            'tariff': old_grid['tariff'], 'emissions': old_grid['emissions']}
    data['electricity'] = [r for r in data['electricity'] if r['name'] != name] + [item]
    letter = openpyxl.utils.get_column_letter(col)
    source('electricity', name, 'price', kf_name, f'25.1!{letter}4:{letter}8763',
           'Arithmetic mean of 8,760 hourly electricity spot prices', 'DKK/MWh', 2026,
           f"annual mean, divide by workbook DKK/EUR ({data['defaults']['dkkPerEur']})")

data['defaults'].update(electricity='KF26 grid spot average - 2030',
                        carbon='AF25 ETS1 allowance - 2030',
                        electrolyser='Alkaline 100 MW - 2030',
                        storage='LOHC -2030', fuelYear=2030, heatYear=2030)
DATA.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
SOURCES.write_text(json.dumps(sources, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print('Updated AF25 fuel and allowance scenarios, KF26 spot-price options, and 2030 defaults')
